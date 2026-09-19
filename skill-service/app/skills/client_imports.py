import csv
import hashlib
import hmac
import io
import json
import re
import secrets
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Literal
from uuid import uuid4

from openpyxl import load_workbook
from pydantic import (
    BaseModel,
    ConfigDict,
    Field,
    ValidationError,
    field_validator,
)
from sqlalchemy import bindparam, text

from app.database.connection import database_transaction
from app.database.registry import get_database_config
from app.database.tables import get_table_definition


MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 500
CONFIRMATION_MINUTES = 10

SUPPORTED_COLUMNS = {
    "client_code",
    "name",
    "business_type",
    "phone",
    "email",
    "city",
    "source",
    "status",
    "notes",
}

REQUIRED_COLUMNS = {
    "client_code",
    "name",
}


class InvalidClientImportFileError(Exception):
    pass


class ClientImportConfirmationError(Exception):
    pass


class ConfirmClientImportParameters(BaseModel):
    model_config = ConfigDict(extra="forbid")

    database_id: str = Field(
        min_length=1,
        max_length=50,
        pattern=r"^[a-z][a-z0-9_]*$",
    )
    action_id: str = Field(
        pattern=(
            r"^[0-9a-f]{8}-[0-9a-f]{4}-"
            r"[0-9a-f]{4}-[0-9a-f]{4}-"
            r"[0-9a-f]{12}$"
        ),
    )
    confirmation_token: str = Field(
        min_length=32,
        max_length=100,
    )


class ClientImportRow(BaseModel):
    model_config = ConfigDict(extra="forbid")

    client_code: str = Field(
        min_length=1,
        max_length=50,
        pattern=r"^[A-Za-z0-9][A-Za-z0-9_-]*$",
    )
    name: str = Field(min_length=1, max_length=150)
    business_type: str | None = Field(
        default=None,
        min_length=1,
        max_length=100,
    )
    phone: str | None = Field(
        default=None,
        min_length=1,
        max_length=30,
    )
    email: str | None = Field(
        default=None,
        min_length=3,
        max_length=255,
    )
    city: str | None = Field(
        default=None,
        min_length=1,
        max_length=100,
    )
    source: str | None = Field(
        default=None,
        min_length=1,
        max_length=100,
    )
    status: Literal[
        "prospect",
        "lead",
        "contacted",
        "follow_up",
        "qualified",
        "won",
        "lost",
    ] = "prospect"
    notes: str | None = Field(default=None)

    @field_validator("*", mode="before")
    @classmethod
    def normalize_values(
        cls,
        value: Any,
    ) -> Any:
        if value is None:
            return None

        if isinstance(value, float) and value.is_integer():
            value = int(value)

        if isinstance(value, (int, float)):
            return str(value)

        if isinstance(value, str):
            normalized = value.strip()
            return normalized or None

        return value


def normalize_header(value: Any) -> str:
    if value is None:
        return ""

    normalized = str(value).strip().lower()
    normalized = re.sub(r"[\s-]+", "_", normalized)

    return normalized


def normalize_cell(value: Any) -> Any:
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.isoformat()

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    if isinstance(value, (int, float)):
        return str(value)

    if isinstance(value, str):
        normalized = value.strip()
        return normalized or None

    return str(value).strip() or None


def validate_headers(headers: list[str]) -> None:
    if not headers or all(not header for header in headers):
        raise InvalidClientImportFileError(
            "Header file tidak ditemukan"
        )

    if len(headers) != len(set(headers)):
        raise InvalidClientImportFileError(
            "Header file memiliki kolom duplikat"
        )

    missing = REQUIRED_COLUMNS.difference(headers)

    if missing:
        raise InvalidClientImportFileError(
            "Kolom wajib tidak tersedia: "
            + ", ".join(sorted(missing))
        )

    unsupported = set(headers).difference(SUPPORTED_COLUMNS)

    if unsupported:
        raise InvalidClientImportFileError(
            "Kolom tidak didukung: "
            + ", ".join(sorted(unsupported))
        )


def parse_csv_file(
    file_content: bytes,
) -> list[dict[str, Any]]:
    try:
        decoded = file_content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise InvalidClientImportFileError(
            "CSV harus menggunakan encoding UTF-8"
        ) from error

    try:
        sample = decoded[:4096]
        dialect = csv.Sniffer().sniff(
            sample,
            delimiters=",;",
        )
    except csv.Error:
        dialect = csv.excel

    reader = csv.reader(
        io.StringIO(decoded),
        dialect=dialect,
    )

    try:
        raw_headers = next(reader)
    except StopIteration as error:
        raise InvalidClientImportFileError(
            "File CSV kosong"
        ) from error

    headers = [
        normalize_header(header)
        for header in raw_headers
    ]
    validate_headers(headers)

    rows: list[dict[str, Any]] = []

    for values in reader:
        if not any(
            normalize_cell(value) is not None
            for value in values
        ):
            continue

        if len(values) > len(headers):
            raise InvalidClientImportFileError(
                "Jumlah kolom CSV tidak konsisten"
            )

        padded_values = values + [None] * (
            len(headers) - len(values)
        )

        rows.append(
            {
                header: normalize_cell(value)
                for header, value in zip(
                    headers,
                    padded_values,
                    strict=True,
                )
            }
        )

        if len(rows) > MAX_IMPORT_ROWS:
            raise InvalidClientImportFileError(
                f"File melebihi {MAX_IMPORT_ROWS} baris"
            )

    return rows


def parse_xlsx_file(
    file_content: bytes,
) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(
            io.BytesIO(file_content),
            read_only=True,
            data_only=True,
        )
    except Exception as error:
        raise InvalidClientImportFileError(
            "File XLSX tidak dapat dibaca"
        ) from error

    try:
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)

        try:
            raw_headers = next(iterator)
        except StopIteration as error:
            raise InvalidClientImportFileError(
                "File XLSX kosong"
            ) from error

        headers = [
            normalize_header(header)
            for header in raw_headers
        ]
        validate_headers(headers)

        rows: list[dict[str, Any]] = []

        for values in iterator:
            normalized_values = [
                normalize_cell(value)
                for value in values[:len(headers)]
            ]

            if not any(
                value is not None
                for value in normalized_values
            ):
                continue

            normalized_values += [None] * (
                len(headers) - len(normalized_values)
            )

            rows.append(
                {
                    header: value
                    for header, value in zip(
                        headers,
                        normalized_values,
                        strict=True,
                    )
                }
            )

            if len(rows) > MAX_IMPORT_ROWS:
                raise InvalidClientImportFileError(
                    f"File melebihi "
                    f"{MAX_IMPORT_ROWS} baris"
                )

        return rows
    finally:
        workbook.close()


def parse_client_file(
    file_name: str,
    file_content: bytes,
) -> list[dict[str, Any]]:
    if not file_content:
        raise InvalidClientImportFileError(
            "File tidak memiliki isi"
        )

    if len(file_content) > MAX_FILE_SIZE:
        raise InvalidClientImportFileError(
            "Ukuran file melebihi 5 MB"
        )

    extension = Path(file_name).suffix.lower()

    if extension == ".csv":
        return parse_csv_file(file_content)

    if extension == ".xlsx":
        return parse_xlsx_file(file_content)

    raise InvalidClientImportFileError(
        "Hanya file .xlsx dan .csv yang didukung"
    )


def get_existing_client_codes(
    connection,
    table_name: str,
    client_codes: list[str],
) -> set[str]:
    if not client_codes:
        return set()

    statement = text(
        f"""
        SELECT `client_code`
        FROM `{table_name}`
        WHERE `client_code` IN :client_codes
        """
    ).bindparams(
        bindparam(
            "client_codes",
            expanding=True,
        )
    )

    rows = connection.execute(
        statement,
        {
            "client_codes": client_codes,
        },
    ).scalars().all()

    return set(rows)


def preview_client_import(
    role: str,
    actor_id: str,
    database_id: str,
    file_name: str,
    mime_type: str,
    file_content: bytes,
) -> dict[str, Any]:
    if role != "superadmin":
        raise PermissionError

    if not re.fullmatch(r"^[a-f0-9]{64}$", actor_id):
        raise PermissionError

    get_database_config(database_id, role)

    table = get_table_definition(
        database_id,
        "clients",
        role,
    )

    raw_rows = parse_client_file(
        file_name,
        file_content,
    )

    valid_rows: list[ClientImportRow] = []
    invalid_rows: list[dict[str, Any]] = []
    duplicate_file_codes: list[str] = []
    seen_codes: set[str] = set()

    for row_number, raw_row in enumerate(
        raw_rows,
        start=2,
    ):
        try:
            validated = ClientImportRow.model_validate(
                raw_row
            )
        except ValidationError as error:
            invalid_rows.append(
                {
                    "row": row_number,
                    "errors": [
                        {
                            "field": ".".join(
                                str(part)
                                for part in item["loc"]
                            ),
                            "message": item["msg"],
                        }
                        for item in error.errors()
                    ],
                }
            )
            continue

        if validated.client_code in seen_codes:
            duplicate_file_codes.append(
                validated.client_code
            )
            continue

        seen_codes.add(validated.client_code)
        valid_rows.append(validated)

    file_hash = hashlib.sha256(file_content).hexdigest()

    with database_transaction(
        database_id,
        role,
    ) as connection:
        existing_codes = get_existing_client_codes(
            connection,
            table.physical_name,
            [
                row.client_code
                for row in valid_rows
            ],
        )

        new_rows = [
            row
            for row in valid_rows
            if row.client_code not in existing_codes
        ]

        action_id = None
        confirmation_token = None
        expires_at = None

        if new_rows:
            action_id = str(uuid4())
            confirmation_token = secrets.token_urlsafe(32)
            confirmation_token_hash = hashlib.sha256(
                confirmation_token.encode("utf-8")
            ).hexdigest()

            now = datetime.now(timezone.utc).replace(
                tzinfo=None
            )
            expires_at = now + timedelta(
                minutes=CONFIRMATION_MINUTES
            )

            payload = {
                "file_name": Path(file_name).name,
                "mime_type": mime_type,
                "content_hash": file_hash,
                "rows": [
                    row.model_dump()
                    for row in new_rows
                ],
            }

            connection.execute(
                text(
                    """
                    INSERT INTO `pending_actions` (
                        `action_id`,
                        `requested_by`,
                        `action_type`,
                        `database_id`,
                        `resource_id`,
                        `payload_json`,
                        `confirmation_token_hash`,
                        `status`,
                        `expires_at`
                    )
                    VALUES (
                        :action_id,
                        :requested_by,
                        'import_clients',
                        :database_id,
                        :resource_id,
                        :payload_json,
                        :confirmation_token_hash,
                        'pending',
                        :expires_at
                    )
                    """
                ),
                {
                    "action_id": action_id,
                    "requested_by": actor_id,
                    "database_id": database_id,
                    "resource_id": file_hash,
                    "payload_json": json.dumps(payload),
                    "confirmation_token_hash":
                        confirmation_token_hash,
                    "expires_at": expires_at,
                },
            )

    return {
        "database_id": database_id,
        "file": {
            "name": Path(file_name).name,
            "mime_type": mime_type,
            "size": len(file_content),
            "content_hash": file_hash,
        },
        "summary": {
            "total_rows": len(raw_rows),
            "new_rows": len(new_rows),
            "existing_rows": len(existing_codes),
            "duplicate_file_rows":
                len(duplicate_file_codes),
            "invalid_rows": len(invalid_rows),
        },
        "existing_client_codes": sorted(existing_codes),
        "duplicate_file_codes": duplicate_file_codes,
        "invalid_row_details": invalid_rows[:20],
        "action_id": action_id,
        "confirmation_token": confirmation_token,
        "expires_at": (
            expires_at.isoformat()
            if expires_at
            else None
        ),
    }


def confirm_client_import(
    role: str,
    actor_id: str,
    database_id: str,
    action_id: str,
    confirmation_token: str,
) -> dict[str, Any]:
    if role != "superadmin":
        raise PermissionError

    if not re.fullmatch(r"^[a-f0-9]{64}$", actor_id):
        raise PermissionError

    get_database_config(database_id, role)

    table = get_table_definition(
        database_id,
        "clients",
        role,
    )

    with database_transaction(
        database_id,
        role,
    ) as connection:
        pending = connection.execute(
            text(
                """
                SELECT
                    `id`,
                    `requested_by`,
                    `payload_json`,
                    `confirmation_token_hash`,
                    `status`,
                    `expires_at`
                FROM `pending_actions`
                WHERE `action_id` = :action_id
                  AND `action_type` = 'import_clients'
                  AND `database_id` = :database_id
                LIMIT 1
                FOR UPDATE
                """
            ),
            {
                "action_id": action_id,
                "database_id": database_id,
            },
        ).mappings().first()

        now = datetime.now(timezone.utc).replace(
            tzinfo=None
        )

        if (
            pending is None
            or pending["requested_by"] != actor_id
            or pending["status"] != "pending"
            or pending["expires_at"] <= now
        ):
            raise ClientImportConfirmationError

        supplied_hash = hashlib.sha256(
            confirmation_token.encode("utf-8")
        ).hexdigest()

        if not hmac.compare_digest(
            supplied_hash,
            pending["confirmation_token_hash"],
        ):
            raise ClientImportConfirmationError

        payload = json.loads(
            pending["payload_json"]
        )
        rows = payload.get("rows", [])

        validated_rows = [
            ClientImportRow.model_validate(row)
            for row in rows
        ]

        existing_codes = get_existing_client_codes(
            connection,
            table.physical_name,
            [
                row.client_code
                for row in validated_rows
            ],
        )

        rows_to_insert = [
            row.model_dump()
            for row in validated_rows
            if row.client_code not in existing_codes
        ]

        if rows_to_insert:
            connection.execute(
                text(
                    f"""
                    INSERT INTO `{table.physical_name}` (
                        `client_code`,
                        `name`,
                        `business_type`,
                        `phone`,
                        `email`,
                        `city`,
                        `source`,
                        `status`,
                        `notes`
                    )
                    VALUES (
                        :client_code,
                        :name,
                        :business_type,
                        :phone,
                        :email,
                        :city,
                        :source,
                        :status,
                        :notes
                    )
                    """
                ),
                rows_to_insert,
            )

        connection.execute(
            text(
                """
                UPDATE `pending_actions`
                SET
                    `status` = 'confirmed',
                    `confirmed_at` = :confirmed_at
                WHERE `id` = :pending_id
                """
            ),
            {
                "confirmed_at": now,
                "pending_id": pending["id"],
            },
        )

    return {
        "database_id": database_id,
        "imported": True,
        "file_name": payload["file_name"],
        "imported_count": len(rows_to_insert),
        "skipped_existing_count": len(existing_codes),
        "imported_client_codes": [
            row["client_code"]
            for row in rows_to_insert
        ],
    }