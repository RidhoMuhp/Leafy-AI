import csv
import hashlib
import os
import re
from datetime import datetime, timezone, timedelta
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any
from uuid import uuid4
from zipfile import BadZipFile, ZipFile


import pymupdf
import cv2
import numpy as np
import pytesseract
from docx import Document
from openpyxl import load_workbook
from PIL import (
    Image,
    ImageFilter,
    ImageOps,
    UnidentifiedImageError,
)
from pypdf import PdfReader
from sqlalchemy import text

from app.database.connection import database_transaction
from app.database.registry import get_database_config
from app.database.tables import get_table_definition


OCR_DIGIT_REPLACEMENTS = str.maketrans(
    {
        "O": "0",
        "o": "0",
        "Q": "0",
        "q": "0",
        "D": "0",
        "E": "0",
        "e": "0",
        "I": "1",
        "l": "1",
        "|": "1",
        "S": "5",
        "s": "5",
        "B": "8",
    }
)
MAX_DOCUMENT_FILE_SIZE = 15 * 1024 * 1024
MAX_PDF_PAGES = 50
MAX_EXTRACTED_CHARACTERS = 1_000_000
MAX_EXCEL_SHEETS = 20
MAX_EXCEL_ROWS_PER_SHEET = 10_000
CHUNK_SIZE = 3000
CHUNK_OVERLAP = 300

ALLOWED_EXTENSIONS = {
    ".pdf": "application/pdf",
    ".docx": (
        "application/vnd.openxmlformats-officedocument."
        "wordprocessingml.document"
    ),
    ".xlsx": (
        "application/vnd.openxmlformats-officedocument."
        "spreadsheetml.sheet"
    ),
    ".csv": "text/csv",
    ".txt": "text/plain",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
}

WITA_TIMEZONE = timezone(
    timedelta(hours=8),
    name="WITA",
)

class InvalidDocumentFileError(Exception):
    pass


class DuplicateDocumentError(Exception):
    pass


class DocumentExtractionError(Exception):
    pass


def configure_tesseract() -> None:
    configured_path = os.getenv("TESSERACT_CMD")

    if configured_path:
        pytesseract.pytesseract.tesseract_cmd = (
            configured_path
        )


def sanitize_file_name(file_name: str) -> str:
    original_name = Path(file_name).name.strip()

    if not original_name:
        raise InvalidDocumentFileError(
            "Nama dokumen tidak valid"
        )

    if len(original_name) > 255:
        raise InvalidDocumentFileError(
            "Nama dokumen terlalu panjang"
        )

    stem = Path(original_name).stem
    suffix = Path(original_name).suffix.lower()

    safe_stem = re.sub(
        r"[^A-Za-z0-9._-]+",
        "-",
        stem,
    ).strip("._-")

    if not safe_stem:
        safe_stem = "document"

    return f"{safe_stem[:180]}{suffix}"


def validate_zip_structure(
    file_content: bytes,
    extension: str,
) -> None:
    try:
        with ZipFile(BytesIO(file_content)) as archive:
            names = set(archive.namelist())
    except BadZipFile as error:
        raise InvalidDocumentFileError(
            "Struktur file ZIP Office tidak valid"
        ) from error

    required_name = (
        "word/document.xml"
        if extension == ".docx"
        else "xl/workbook.xml"
    )

    if required_name not in names:
        raise InvalidDocumentFileError(
            "Isi file tidak sesuai dengan ekstensi"
        )


def validate_document_file(
    file_name: str,
    file_content: bytes,
) -> tuple[str, str]:
    if not file_content:
        raise InvalidDocumentFileError(
            "Dokumen kosong"
        )

    if len(file_content) > MAX_DOCUMENT_FILE_SIZE:
        raise InvalidDocumentFileError(
            "Ukuran dokumen melebihi 15 MB"
        )

    extension = Path(file_name).suffix.lower()
    mime_type = ALLOWED_EXTENSIONS.get(extension)

    if mime_type is None:
        raise InvalidDocumentFileError(
            "Jenis dokumen tidak didukung"
        )

    if extension == ".pdf":
        if not file_content.startswith(b"%PDF-"):
            raise InvalidDocumentFileError(
                "Signature PDF tidak valid"
            )

    elif extension in {".docx", ".xlsx"}:
        validate_zip_structure(
            file_content,
            extension,
        )

    elif extension == ".png":
        if not file_content.startswith(
            b"\x89PNG\r\n\x1a\n"
        ):
            raise InvalidDocumentFileError(
                "Signature PNG tidak valid"
            )

    elif extension in {".jpg", ".jpeg"}:
        if not file_content.startswith(b"\xff\xd8\xff"):
            raise InvalidDocumentFileError(
                "Signature JPEG tidak valid"
            )

    elif b"\x00" in file_content[:4096]:
        raise InvalidDocumentFileError(
            "File teks mengandung data biner"
        )

    return extension, mime_type


def decode_text_file(file_content: bytes) -> str:
    for encoding in (
        "utf-8-sig",
        "utf-8",
        "cp1252",
        "latin-1",
    ):
        try:
            return file_content.decode(encoding)
        except UnicodeDecodeError:
            continue

    raise DocumentExtractionError(
        "Encoding file teks tidak dapat dibaca"
    )


def extract_text_file(
    file_content: bytes,
) -> tuple[str, str]:
    return decode_text_file(file_content), "text"


def extract_csv_file(
    file_content: bytes,
) -> tuple[str, str]:
    decoded = decode_text_file(file_content)

    try:
        rows = csv.reader(StringIO(decoded))
        lines = []

        for row_number, row in enumerate(
            rows,
            start=1,
        ):
            cleaned = [
                str(value).strip()
                for value in row
            ]

            lines.append(
                f"Baris {row_number}: "
                + " | ".join(cleaned)
            )

            if row_number >= MAX_EXCEL_ROWS_PER_SHEET:
                break

        return "\n".join(lines), "csv"

    except csv.Error as error:
        raise DocumentExtractionError(
            "CSV tidak dapat dibaca"
        ) from error


def extract_docx_file(
    file_content: bytes,
) -> tuple[str, str]:
    try:
        document = Document(BytesIO(file_content))
    except Exception as error:
        raise DocumentExtractionError(
            "DOCX tidak dapat dibaca"
        ) from error

    sections = []

    for paragraph in document.paragraphs:
        content = paragraph.text.strip()

        if content:
            sections.append(content)

    for table_number, table in enumerate(
        document.tables,
        start=1,
    ):
        sections.append(f"[Tabel {table_number}]")

        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
            ]
            sections.append(" | ".join(values))

    return "\n".join(sections), "docx"


def extract_xlsx_file(
    file_content: bytes,
) -> tuple[str, str]:
    try:
        workbook = load_workbook(
            BytesIO(file_content),
            read_only=True,
            data_only=True,
        )
    except Exception as error:
        raise DocumentExtractionError(
            "XLSX tidak dapat dibaca"
        ) from error

    sections = []

    try:
        for sheet_number, worksheet in enumerate(
            workbook.worksheets,
            start=1,
        ):
            if sheet_number > MAX_EXCEL_SHEETS:
                break

            sections.append(
                f"[Sheet: {worksheet.title}]"
            )

            for row_number, row in enumerate(
                worksheet.iter_rows(values_only=True),
                start=1,
            ):
                if row_number > MAX_EXCEL_ROWS_PER_SHEET:
                    break

                values = [
                    "" if value is None else str(value)
                    for value in row
                ]

                if any(value.strip() for value in values):
                    sections.append(
                        f"Baris {row_number}: "
                        + " | ".join(values)
                    )
    finally:
        workbook.close()

    return "\n".join(sections), "xlsx"

def merge_ocr_results(
    results: list[str],
) -> str:
    lines = []
    seen = set()

    for result in results:
        for raw_line in result.splitlines():
            line = raw_line.strip()

            if not line:
                continue

            comparison = re.sub(
                r"[^a-z0-9]+",
                "",
                line.lower(),
            )

            if (
                not comparison
                or comparison in seen
            ):
                continue

            seen.add(comparison)
            lines.append(line)

    return "\n".join(lines)

def prepare_ocr_image(
    image: Image.Image,
) -> Image.Image:
    grayscale = ImageOps.grayscale(
        image
    )

    enhanced = ImageOps.autocontrast(
        grayscale,
        cutoff=1,
    )

    width, height = enhanced.size
    scale = 2

    if width * scale > 3000:
        scale = max(
            1,
            3000 // max(width, 1),
        )

    if scale > 1:
        enhanced = enhanced.resize(
            (
                width * scale,
                height * scale,
            ),
            Image.Resampling.LANCZOS,
        )

    return enhanced.filter(
        ImageFilter.SHARPEN
    )

def create_ocr_variants(
    image: Image.Image,
) -> list[Image.Image]:
    binary_normal = image.point(
        lambda pixel:
            255 if pixel > 150 else 0
    )

    binary_bright = image.point(
        lambda pixel:
            255 if pixel > 180 else 0
    )

    binary_inverted = ImageOps.invert(
        binary_bright
    )

    return [
        image,
        binary_normal,
        binary_inverted,
    ]


def pil_to_grayscale_array(
    image: Image.Image,
) -> np.ndarray:
    rgb = np.array(
        ImageOps.exif_transpose(image).convert("RGB")
    )
    return cv2.cvtColor(
        rgb,
        cv2.COLOR_RGB2GRAY,
    )


def crop_image_content(
    grayscale: np.ndarray,
) -> np.ndarray:
    blurred = cv2.GaussianBlur(
        grayscale,
        (5, 5),
        0,
    )
    _, mask = cv2.threshold(
        blurred,
        0,
        255,
        cv2.THRESH_BINARY_INV
        + cv2.THRESH_OTSU,
    )
    points = cv2.findNonZero(mask)

    if points is None:
        return grayscale

    x, y, width, height = cv2.boundingRect(points)
    image_height, image_width = grayscale.shape
    padding = max(
        12,
        int(min(image_width, image_height) * 0.015),
    )
    left = max(0, x - padding)
    top = max(0, y - padding)
    right = min(image_width, x + width + padding)
    bottom = min(image_height, y + height + padding)

    cropped = grayscale[top:bottom, left:right]
    return cropped if cropped.size else grayscale


def deskew_grayscale(
    grayscale: np.ndarray,
) -> np.ndarray:
    inverted = cv2.bitwise_not(grayscale)
    _, thresholded = cv2.threshold(
        inverted,
        0,
        255,
        cv2.THRESH_BINARY
        + cv2.THRESH_OTSU,
    )
    coordinates = np.column_stack(
        np.where(thresholded > 0)
    )

    if len(coordinates) < 30:
        return grayscale

    angle = cv2.minAreaRect(coordinates)[-1]
    angle = -(90 + angle) if angle < -45 else -angle

    if abs(angle) < 0.15 or abs(angle) > 12:
        return grayscale

    height, width = grayscale.shape
    center = (width / 2, height / 2)
    matrix = cv2.getRotationMatrix2D(
        center,
        angle,
        1.0,
    )
    return cv2.warpAffine(
        grayscale,
        matrix,
        (width, height),
        flags=cv2.INTER_CUBIC,
        borderMode=cv2.BORDER_REPLICATE,
    )


def resize_for_ocr(
    grayscale: np.ndarray,
) -> np.ndarray:
    height, width = grayscale.shape
    longest_side = max(height, width)

    if longest_side < 1400:
        scale = min(3.0, 1800 / max(longest_side, 1))
    elif longest_side < 2400:
        scale = 1.5
    else:
        scale = 1.0

    if scale <= 1.0:
        return grayscale

    return cv2.resize(
        grayscale,
        None,
        fx=scale,
        fy=scale,
        interpolation=cv2.INTER_CUBIC,
    )


def create_advanced_ocr_variants(
    image: Image.Image,
) -> list[tuple[str, Image.Image]]:
    grayscale = pil_to_grayscale_array(image)
    grayscale = crop_image_content(grayscale)
    grayscale = deskew_grayscale(grayscale)
    grayscale = resize_for_ocr(grayscale)

    denoised = cv2.fastNlMeansDenoising(
        grayscale,
        None,
        h=9,
        templateWindowSize=7,
        searchWindowSize=21,
    )
    clahe = cv2.createCLAHE(
        clipLimit=2.2,
        tileGridSize=(8, 8),
    ).apply(denoised)
    sharpened = cv2.addWeighted(
        clahe,
        1.45,
        cv2.GaussianBlur(clahe, (0, 0), 1.2),
        -0.45,
        0,
    )
    adaptive = cv2.adaptiveThreshold(
        sharpened,
        255,
        cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
        cv2.THRESH_BINARY,
        35,
        11,
    )
    _, otsu = cv2.threshold(
        sharpened,
        0,
        255,
        cv2.THRESH_BINARY + cv2.THRESH_OTSU,
    )

    arrays = [
        ("grayscale", grayscale),
        ("clahe", sharpened),
        ("adaptive", adaptive),
        ("otsu", otsu),
    ]
    return [
        (name, Image.fromarray(array))
        for name, array in arrays
    ]


def run_scored_tesseract(
    image: Image.Image,
    psm: int,
) -> tuple[str, float]:
    config = (
        "--oem 3 "
        f"--psm {psm} "
        "-c preserve_interword_spaces=1"
    )
    data = pytesseract.image_to_data(
        image,
        config=config,
        output_type=pytesseract.Output.DICT,
    )
    words: list[str] = []
    confidences: list[float] = []
    current_line: tuple[int, int, int] | None = None
    lines: list[list[str]] = []

    for index, raw_word in enumerate(data["text"]):
        word = str(raw_word).strip()
        if not word:
            continue

        try:
            confidence = float(data["conf"][index])
        except (TypeError, ValueError):
            confidence = -1

        if confidence < 10:
            continue

        line_key = (
            int(data["block_num"][index]),
            int(data["par_num"][index]),
            int(data["line_num"][index]),
        )
        if line_key != current_line:
            lines.append([])
            current_line = line_key

        lines[-1].append(word)
        words.append(word)
        confidences.append(confidence)

    content = "\n".join(
        " ".join(line).strip()
        for line in lines
        if line
    )
    average_confidence = (
        sum(confidences) / len(confidences)
        if confidences
        else 0.0
    )
    return content, average_confidence

def extract_positioned_words(
    image: Image.Image,
) -> list[dict[str, Any]]:
    data = pytesseract.image_to_data(
        image,
        config=(
            "--oem 3 "
            "--psm 11 "
            "-c preserve_interword_spaces=1"
        ),
        output_type=(
            pytesseract.Output.DICT
        ),
    )

    words = []

    for index, raw_text in enumerate(
        data["text"]
    ):
        word = str(raw_text).strip()

        if not word:
            continue

        try:
            confidence = float(
                data["conf"][index]
            )
        except (
            TypeError,
            ValueError,
        ):
            confidence = -1

        if confidence < 15:
            continue

        words.append(
            {
                "text": word,
                "confidence": confidence,
                "left": int(
                    data["left"][index]
                ),
                "top": int(
                    data["top"][index]
                ),
                "width": int(
                    data["width"][index]
                ),
                "height": int(
                    data["height"][index]
                ),
                "block": int(
                    data["block_num"][index]
                ),
                "paragraph": int(
                    data["par_num"][index]
                ),
                "line": int(
                    data["line_num"][index]
                ),
            }
        )

    return words

def reconstruct_ocr_lines(
    words: list[dict[str, Any]],
) -> list[str]:
    grouped: dict[
        tuple[int, int, int],
        list[dict[str, Any]],
    ] = {}

    for word in words:
        key = (
            word["block"],
            word["paragraph"],
            word["line"],
        )

        grouped.setdefault(
            key,
            [],
        ).append(word)

    ordered_groups = sorted(
        grouped.values(),
        key=lambda group: (
            min(
                item["top"]
                for item in group
            ),
            min(
                item["left"]
                for item in group
            ),
        ),
    )

    lines = []

    for group in ordered_groups:
        ordered_words = sorted(
            group,
            key=lambda item:
                item["left"],
        )

        line = " ".join(
            item["text"]
            for item in ordered_words
        ).strip()

        if line:
            lines.append(line)

    return lines

def create_transaction_crops(
    image: Image.Image,
    words: list[dict[str, Any]],
) -> list[Image.Image]:
    width, height = image.size
    crops = []

    # Bagian atas sering memuat status
    # transaksi dan nominal.
    top_crop_height = min(
        height,
        max(300, int(height * 0.4)),
    )

    crops.append(
        image.crop(
            (
                0,
                0,
                width,
                top_crop_height,
            )
        )
    )

    anchors = (
        "transfer",
        "pembayaran",
        "transaksi",
        "nominal",
        "jumlah",
        "total",
    )

    matching_words = [
        word
        for word in words
        if any(
            anchor in word[
                "text"
            ].lower()
            for anchor in anchors
        )
    ]

    if matching_words:
        anchor_top = min(
            word["top"]
            for word in matching_words
        )

        anchor_bottom = max(
            word["top"]
            + word["height"]
            for word in matching_words
        )

        crop_top = max(
            0,
            anchor_top - 120,
        )

        crop_bottom = min(
            height,
            anchor_bottom + 400,
        )

        if crop_bottom > crop_top:
            crops.append(
                image.crop(
                    (
                        0,
                        crop_top,
                        width,
                        crop_bottom,
                    )
                )
            )

    return crops

def run_tesseract(
    image: Image.Image,
    psm: int,
) -> str:
    return pytesseract.image_to_string(
        image,
        config=(
            "--oem 3 "
            f"--psm {psm} "
            "-c preserve_interword_spaces=1"
        ),
    )

def ocr_image(
    image: Image.Image,
) -> str:
    configure_tesseract()

    results: list[tuple[float, str]] = []

    try:
        variants = create_advanced_ocr_variants(
            image
        )
        variant_scores: dict[str, float] = {}

        psm_by_variant = {
            "grayscale": (4, 6, 11),
            "clahe": (4, 6, 12),
            "adaptive": (6, 11),
            "otsu": (6, 11),
        }

        for variant_name, variant in variants:
            for psm in psm_by_variant[variant_name]:
                content, confidence = run_scored_tesseract(
                    variant,
                    psm,
                )
                if content.strip():
                    results.append((confidence, content))
                    variant_scores[variant_name] = max(
                        confidence,
                        variant_scores.get(variant_name, 0.0),
                    )

        best_variant = max(
            variants,
            key=lambda item: variant_scores.get(
                item[0],
                0.0,
            ),
        )[1]
        positioned_words = extract_positioned_words(
            best_variant
        )
        positioned_lines = reconstruct_ocr_lines(
            positioned_words
        )
        if positioned_lines:
            results.append(
                (70.0, "\n".join(positioned_lines))
            )

        transaction_crops = create_transaction_crops(
            best_variant,
            positioned_words,
        )
        for crop in transaction_crops:
            for psm in (6, 7, 11):
                content, confidence = run_scored_tesseract(
                    crop,
                    psm,
                )
                if content.strip():
                    results.append(
                        (confidence + 5.0, content)
                    )

    except (
        pytesseract
        .TesseractNotFoundError
    ) as error:
        raise DocumentExtractionError(
            "Tesseract OCR tidak ditemukan"
        ) from error

    except pytesseract.TesseractError as error:
        raise DocumentExtractionError(
            "OCR gagal membaca gambar"
        ) from error

    ordered_results = [
        content
        for _, content in sorted(
            results,
            key=lambda item: item[0],
            reverse=True,
        )
    ]
    merged = merge_ocr_results(ordered_results)
    
    return normalize_ocr_currency(
        merged
    )

def extract_image_file(
    file_content: bytes,
) -> tuple[str, str]:
    try:
        with Image.open(BytesIO(file_content)) as image:
            image.verify()

        with Image.open(BytesIO(file_content)) as image:
            content = ocr_image(image)

    except UnidentifiedImageError as error:
        raise InvalidDocumentFileError(
            "File gambar tidak valid"
        ) from error

    return content, "image_ocr"


def extract_pdf_file(
    file_content: bytes,
) -> tuple[str, str]:
    try:
        reader = PdfReader(BytesIO(file_content))
    except Exception as error:
        raise DocumentExtractionError(
            "PDF tidak dapat dibaca"
        ) from error

    if len(reader.pages) > MAX_PDF_PAGES:
        raise InvalidDocumentFileError(
            "PDF melebihi batas 50 halaman"
        )

    page_contents = []
    used_ocr = False
    pdf_for_ocr = None

    try:
        for page_index, page in enumerate(
            reader.pages,
        ):
            try:
                page_text = (
                    page.extract_text() or ""
                ).strip()
            except Exception:
                page_text = ""

            if len(page_text) < 30:
                used_ocr = True

                if pdf_for_ocr is None:
                    pdf_for_ocr = pymupdf.open(
                        stream=file_content,
                        filetype="pdf",
                    )

                pdf_page = pdf_for_ocr.load_page(
                    page_index
                )

                pixmap = pdf_page.get_pixmap(
                    matrix=pymupdf.Matrix(1.5, 1.5),
                    alpha=False,
                )

                with Image.open(
                    BytesIO(
                        pixmap.tobytes("png")
                    )
                ) as image:
                    page_text = ocr_image(image).strip()

            if page_text:
                page_contents.append(
                    f"[Halaman {page_index + 1}]\n"
                    f"{page_text}"
                )
    finally:
        if pdf_for_ocr is not None:
            pdf_for_ocr.close()

    extraction_method = (
        "pdf_text_ocr"
        if used_ocr
        else "pdf_text"
    )

    return "\n\n".join(page_contents), extraction_method


def normalize_extracted_text(content: str) -> str:
    normalized = content.replace(
        "\r\n",
        "\n",
    ).replace(
        "\r",
        "\n",
    )

    normalized = re.sub(
        r"[ \t]+\n",
        "\n",
        normalized,
    )

    normalized = re.sub(
        r"\n{3,}",
        "\n\n",
        normalized,
    ).strip()

    if not normalized:
        raise DocumentExtractionError(
            "Tidak ada teks yang dapat diekstrak"
        )

    if len(normalized) > MAX_EXTRACTED_CHARACTERS:
        normalized = normalized[
            :MAX_EXTRACTED_CHARACTERS
        ]

    return normalized


def extract_document_text(
    extension: str,
    file_content: bytes,
) -> tuple[str, str]:
    extractors = {
        ".txt": extract_text_file,
        ".csv": extract_csv_file,
        ".docx": extract_docx_file,
        ".xlsx": extract_xlsx_file,
        ".pdf": extract_pdf_file,
        ".png": extract_image_file,
        ".jpg": extract_image_file,
        ".jpeg": extract_image_file,
    }

    extractor = extractors[extension]
    content, extraction_method = extractor(
        file_content
    )

    return (
        normalize_extracted_text(content),
        extraction_method,
    )

def normalize_ocr_currency(
    content: str,
) -> str:
    pattern = re.compile(
        r"\bRp\s*([0-9A-Za-z|.,]+)",
        re.IGNORECASE,
    )

    def replace_currency(
        match: re.Match,
    ) -> str:
        raw_value = match.group(1)

        if not any(
            character.isdigit()
            for character in raw_value
        ):
            return match.group(0)

        normalized_value = (
            raw_value.translate(
                OCR_DIGIT_REPLACEMENTS
            )
        )

        normalized_value = re.sub(
            r"[^0-9.,]",
            "",
            normalized_value,
        )

        if (
            normalized_value == raw_value
            or not normalized_value
        ):
            return match.group(0)

        return (
            f"Rp{normalized_value} "
            f"[OCR asli: Rp{raw_value}]"
        )

    return pattern.sub(
        replace_currency,
        content,
    )
    
def create_chunks(content: str) -> list[str]:
    chunks = []
    start = 0
    content_length = len(content)

    while start < content_length:
        end = min(
            start + CHUNK_SIZE,
            content_length,
        )

        if end < content_length:
            newline_position = content.rfind(
                "\n",
                start,
                end,
            )

            if newline_position > start + 1000:
                end = newline_position

        chunk = content[start:end].strip()

        if chunk:
            chunks.append(chunk)

        if end >= content_length:
            break

        start = max(
            end - CHUNK_OVERLAP,
            start + 1,
        )

    return chunks

DOCUMENT_TYPE_PREFIXES = {
    "transaction": "TRX",
    "invoice": "INV",
    "stock": "STK",
    "report": "RPT",
    "contract": "CNT",
    "general": "DOC",
}


def detect_document_type(content: str) -> str:
    normalized = content.lower()

    if any(
        marker in normalized
        for marker in (
            "transfer berhasil",
            "pembayaran berhasil",
            "transaksi berhasil",
            "pembayaran diterima",
            "dana masuk",
            "transfer masuk",
        )
    ):
        return "transaction"

    if any(
        marker in normalized
        for marker in (
            "invoice",
            "tagihan",
            "jatuh tempo",
        )
    ):
        return "invoice"

    if "stok" in normalized and any(
        marker in normalized
        for marker in (
            "produk",
            "barang",
            "minimum",
            "jumlah",
        )
    ):
        return "stock"

    if any(
        marker in normalized
        for marker in (
            "laporan",
            "report",
            "ringkasan",
        )
    ):
        return "report"

    if any(
        marker in normalized
        for marker in (
            "kontrak",
            "perjanjian",
            "agreement",
        )
    ):
        return "contract"

    return "general"


def get_message_datetime_wita(
    message_timestamp: str | None,
) -> datetime:
    if message_timestamp:
        normalized = message_timestamp.strip()

        if normalized.endswith("Z"):
            normalized = (
                normalized[:-1] + "+00:00"
            )

        try:
            parsed = datetime.fromisoformat(
                normalized
            )

            if parsed.tzinfo is None:
                parsed = parsed.replace(
                    tzinfo=timezone.utc
                )

            return parsed.astimezone(
                WITA_TIMEZONE
            )
        except ValueError:
            pass

    return datetime.now(WITA_TIMEZONE)

def create_short_code(
    connection: Any,
    documents_table: str,
    document_type: str,
    message_timestamp: str | None,
) -> str:
    prefix = DOCUMENT_TYPE_PREFIXES.get(
        document_type,
        "DOC",
    )
    message_datetime = get_message_datetime_wita(
        message_timestamp
    )
    code_prefix = (
        f"{prefix}-{message_datetime.strftime('%d%m%y')}"
    )

    statement = text(
        f"""
        SELECT `short_code`
        FROM `{documents_table}`
        WHERE `short_code` LIKE :code_pattern
        ORDER BY `short_code` DESC
        LIMIT 1
        """
    )
    latest_code = connection.execute(
        statement,
        {"code_pattern": f"{code_prefix}-%"},
    ).scalar_one_or_none()

    sequence = 1
    if latest_code:
        try:
            sequence = int(latest_code.rsplit("-", 1)[1]) + 1
        except (ValueError, IndexError):
            sequence = 1

    if sequence > 99:
        raise DocumentExtractionError(
            "Batas kode dokumen harian tercapai"
        )

    return f"{code_prefix}-{sequence:02d}"


def ingest_document(
    role: str,
    actor_id: str,
    database_id: str,
    file_name: str,
    supplied_mime_type: str,
    file_content: bytes,
    message_timestamp: str | None = None,
) -> dict[str, Any]:
    if role not in {"admin", "superadmin"}:
        raise PermissionError

    get_database_config(database_id, role)
    documents_table = get_table_definition(
        database_id,
        "knowledge_documents",
        role,
    )
    chunks_table = get_table_definition(
        database_id,
        "knowledge_chunks",
        role,
    )

    safe_name = sanitize_file_name(file_name)
    extension, canonical_mime_type = validate_document_file(
        safe_name,
        file_content,
    )
    content_hash = hashlib.sha256(file_content).hexdigest()
    extracted_text, extraction_method = extract_document_text(
        extension,
        file_content,
    )
    chunks = create_chunks(extracted_text)

    if not chunks:
        raise DocumentExtractionError(
            "Dokumen tidak menghasilkan chunk"
        )

    document_type = detect_document_type(extracted_text)
    document_code = "DOC-" + uuid4().hex[:20].upper()

    duplicate_statement = text(
        f"""
        SELECT
            `id`,
            `document_code`,
            `short_code`,
            `original_name`,
            `status`
        FROM `{documents_table.physical_name}`
        WHERE `content_hash` = :content_hash
        LIMIT 1
        """
    )
    insert_document_statement = text(
        f"""
        INSERT INTO `{documents_table.physical_name}` (
            `document_code`,
            `short_code`,
            `original_name`,
            `safe_name`,
            `mime_type`,
            `document_type`,
            `file_size`,
            `status`,
            `uploaded_by`,
            `content_hash`
        )
        VALUES (
            :document_code,
            :short_code,
            :original_name,
            :safe_name,
            :mime_type,
            :document_type,
            :file_size,
            'processed',
            :uploaded_by,
            :content_hash
        )
        """
    )
    insert_chunk_statement = text(
        f"""
        INSERT INTO `{chunks_table.physical_name}` (
            `document_id`,
            `chunk_index`,
            `content`,
            `token_count`
        )
        VALUES (
            :document_id,
            :chunk_index,
            :content,
            :token_count
        )
        """
    )
    select_document_statement = text(
        f"""
        SELECT
            `id`,
            `document_code`,
            `short_code`,
            `original_name`,
            `safe_name`,
            `mime_type`,
            `document_type`,
            `file_size`,
            `status`,
            `uploaded_by`,
            `created_at`,
            `updated_at`
        FROM `{documents_table.physical_name}`
        WHERE `id` = :document_id
        LIMIT 1
        """
    )

    with database_transaction(database_id, role) as connection:
        duplicate = connection.execute(
            duplicate_statement,
            {"content_hash": content_hash},
        ).mappings().first()

        if duplicate is not None:
            raise DuplicateDocumentError(
                duplicate["document_code"]
            )

        short_code = create_short_code(
            connection,
            documents_table.physical_name,
            document_type,
            message_timestamp,
        )
        result = connection.execute(
            insert_document_statement,
            {
                "document_code": document_code,
                "short_code": short_code,
                "original_name": Path(file_name).name,
                "safe_name": safe_name,
                "mime_type": canonical_mime_type,
                "document_type": document_type,
                "file_size": len(file_content),
                "uploaded_by": actor_id,
                "content_hash": content_hash,
            },
        )
        document_id = result.lastrowid

        connection.execute(
            insert_chunk_statement,
            [
                {
                    "document_id": document_id,
                    "chunk_index": index,
                    "content": chunk,
                    "token_count": max(1, len(chunk) // 4),
                }
                for index, chunk in enumerate(chunks)
            ],
        )
        document = connection.execute(
            select_document_statement,
            {"document_id": document_id},
        ).mappings().one()

    return {
        "database_id": database_id,
        "document": dict(document),
        "supplied_mime_type": supplied_mime_type,
        "extraction_method": extraction_method,
        "character_count": len(extracted_text),
        "chunk_count": len(chunks),
        "preview": extracted_text[:500],
        "analysis_text": extracted_text[:5000],
        "duplicate": False,
        "ingested": True,
    }
        
