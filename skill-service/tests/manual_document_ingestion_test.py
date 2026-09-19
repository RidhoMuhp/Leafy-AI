from io import BytesIO
from pathlib import Path
from datetime import datetime
import sys

import httpx
import pymupdf
from docx import Document as DocxDocument
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from sqlalchemy import bindparam, text


sys.path.insert(
    0,
    str(Path(__file__).resolve().parents[1]),
)

from app.config import get_settings
from app.database.connection import database_transaction
from app.skills.document_ingestion import (
    MAX_DOCUMENT_FILE_SIZE,
)


BASE_URL = "http://127.0.0.1:8000"
DATABASE_ID = "leafy_core"
ROLE = "superadmin"

results: list[dict] = []
created_document_ids: list[int] = []


def add_result(
    name: str,
    expected,
    actual,
    detail: str = "",
) -> None:
    results.append(
        {
            "name": name,
            "expected": expected,
            "actual": actual,
            "pass": expected == actual,
            "detail": detail,
        }
    )


def int_value(value) -> int:
    return int(value or 0)


def upload_document(
    client: httpx.Client,
    file_name: str,
    file_content: bytes,
    mime_type: str,
    role: str = ROLE,
) -> httpx.Response:
    return client.post(
        "/documents/ingest",
        data={
            "role": role,
            "database_id": DATABASE_ID,
        },
        files={
            "file": (
                file_name,
                file_content,
                mime_type,
            ),
        },
    )


def create_txt(marker: str) -> bytes:
    return (
        f"{marker}\n"
        "Dokumen pengujian Leafy AI.\n"
        "Dokumen ini digunakan untuk menguji "
        "ekstraksi teks dan penyimpanan chunk."
    ).encode("utf-8")


def create_csv(marker: str) -> bytes:
    content = (
        "client_code,name,status,notes\n"
        f"DOC-001,Demo Dokumen,lead,{marker}\n"
        "DOC-002,Demo Kedua,prospect,"
        "Data pengujian CSV\n"
    )

    return content.encode("utf-8")


def create_xlsx(marker: str) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Data Reachout"

    worksheet.append(
        [
            "client_code",
            "name",
            "status",
            "notes",
        ]
    )
    worksheet.append(
        [
            "XLSX-001",
            "Demo Spreadsheet",
            "lead",
            marker,
        ]
    )

    second_sheet = workbook.create_sheet(
        "Ringkasan"
    )
    second_sheet.append(
        ["keterangan", "nilai"]
    )
    second_sheet.append(
        ["sumber", "manual test"]
    )

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    return buffer.getvalue()


def create_docx(marker: str) -> bytes:
    document = DocxDocument()

    document.add_heading(
        "Leafy Document Intelligence",
        level=1,
    )

    document.add_paragraph(
        marker
    )

    document.add_paragraph(
        "Dokumen ini menguji pembacaan paragraf "
        "dan tabel DOCX."
    )

    table = document.add_table(
        rows=2,
        cols=2,
    )

    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Status"
    table.cell(1, 1).text = "Valid"

    buffer = BytesIO()
    document.save(buffer)

    return buffer.getvalue()


def create_pdf(marker: str) -> bytes:
    document = pymupdf.open()
    page = document.new_page()

    page.insert_text(
        (72, 72),
        "Leafy Document Intelligence",
        fontsize=18,
    )

    page.insert_text(
        (72, 110),
        marker,
        fontsize=14,
    )

    page.insert_text(
        (72, 145),
        (
            "This PDF contains selectable text "
            "for extraction validation."
        ),
        fontsize=12,
    )

    content = document.tobytes()
    document.close()

    return content


def get_test_font(size: int):
    candidates = [
        "arial.ttf",
        "C:\\Windows\\Fonts\\arial.ttf",
        "C:\\Windows\\Fonts\\calibri.ttf",
    ]

    for candidate in candidates:
        try:
            return ImageFont.truetype(
                candidate,
                size,
            )
        except OSError:
            continue

    return ImageFont.load_default()


def create_ocr_image(
    marker: str,
    image_format: str,
) -> bytes:
    image = Image.new(
        "RGB",
        (1600, 420),
        "white",
    )

    draw = ImageDraw.Draw(image)
    title_font = get_test_font(64)
    body_font = get_test_font(44)

    draw.text(
        (80, 70),
        "LEAFY OCR DOCUMENT",
        fill="black",
        font=title_font,
    )

    draw.text(
        (80, 180),
        marker,
        fill="black",
        font=body_font,
    )

    draw.text(
        (80, 270),
        "BETA VALIDATION 2026",
        fill="black",
        font=body_font,
    )

    buffer = BytesIO()

    save_options = {}

    if image_format.upper() == "JPEG":
        save_options["quality"] = 95

    image.save(
        buffer,
        format=image_format,
        **save_options,
    )

    image.close()

    return buffer.getvalue()


def test_valid_upload(
    client: httpx.Client,
    label: str,
    file_name: str,
    file_content: bytes,
    mime_type: str,
    expected_method: str,
    expected_marker: str,
) -> dict | None:
    response = upload_document(
        client=client,
        file_name=file_name,
        file_content=file_content,
        mime_type=mime_type,
    )

    add_result(
        f"Upload {label}",
        200,
        response.status_code,
        response.text,
    )

    if response.status_code != 200:
        return None

    result = response.json()["result"]
    document = result["document"]

    created_document_ids.append(
        document["id"]
    )

    add_result(
        f"{label} ingested",
        True,
        result.get("ingested"),
        response.text,
    )

    add_result(
        f"{label} extraction method",
        expected_method,
        result.get("extraction_method"),
        response.text,
    )

    add_result(
        f"{label} character count",
        True,
        int_value(
            result.get("character_count")
        ) > 0,
        response.text,
    )

    add_result(
        f"{label} chunk count",
        True,
        int_value(
            result.get("chunk_count")
        ) > 0,
        response.text,
    )

    preview = str(
        result.get("preview", "")
    ).upper()

    add_result(
        f"{label} preview marker",
        True,
        expected_marker.upper() in preview,
        response.text,
    )

    return result


def inspect_database() -> None:
    if not created_document_ids:
        add_result(
            "Document metadata stored",
            8,
            0,
        )
        add_result(
            "Documents processed",
            8,
            0,
        )
        add_result(
            "Documents have chunks",
            8,
            0,
        )
        add_result(
            "Chunk rows stored",
            True,
            False,
        )
        return

    document_statement = (
        text(
            """
            SELECT
                COUNT(*) AS document_count,
                COALESCE(
                    SUM(
                        CASE
                            WHEN status = 'processed'
                            THEN 1
                            ELSE 0
                        END
                    ),
                    0
                ) AS processed_count
            FROM knowledge_documents
            WHERE id IN :document_ids
            """
        )
        .bindparams(
            bindparam(
                "document_ids",
                expanding=True,
            )
        )
    )

    chunk_statement = (
        text(
            """
            SELECT
                COUNT(*) AS chunk_count,
                COUNT(DISTINCT document_id)
                    AS document_count
            FROM knowledge_chunks
            WHERE document_id IN :document_ids
            """
        )
        .bindparams(
            bindparam(
                "document_ids",
                expanding=True,
            )
        )
    )

    with database_transaction(
        DATABASE_ID,
        ROLE,
    ) as connection:
        document_result = connection.execute(
            document_statement,
            {
                "document_ids":
                    created_document_ids,
            },
        ).mappings().one()

        chunk_result = connection.execute(
            chunk_statement,
            {
                "document_ids":
                    created_document_ids,
            },
        ).mappings().one()

    expected_count = len(
        created_document_ids
    )

    add_result(
        "Document metadata stored",
        expected_count,
        int_value(
            document_result["document_count"]
        ),
    )

    add_result(
        "Documents processed",
        expected_count,
        int_value(
            document_result["processed_count"]
        ),
    )

    add_result(
        "Documents have chunks",
        expected_count,
        int_value(
            chunk_result["document_count"]
        ),
    )

    add_result(
        "Chunk rows stored",
        True,
        int_value(
            chunk_result["chunk_count"]
        ) >= expected_count,
    )


def cleanup() -> tuple[int, int]:
    if not created_document_ids:
        return 0, 0

    delete_chunks = (
        text(
            """
            DELETE FROM knowledge_chunks
            WHERE document_id IN :document_ids
            """
        )
        .bindparams(
            bindparam(
                "document_ids",
                expanding=True,
            )
        )
    )

    delete_documents = (
        text(
            """
            DELETE FROM knowledge_documents
            WHERE id IN :document_ids
            """
        )
        .bindparams(
            bindparam(
                "document_ids",
                expanding=True,
            )
        )
    )

    count_documents = (
        text(
            """
            SELECT COUNT(*)
            FROM knowledge_documents
            WHERE id IN :document_ids
            """
        )
        .bindparams(
            bindparam(
                "document_ids",
                expanding=True,
            )
        )
    )

    count_chunks = (
        text(
            """
            SELECT COUNT(*)
            FROM knowledge_chunks
            WHERE document_id IN :document_ids
            """
        )
        .bindparams(
            bindparam(
                "document_ids",
                expanding=True,
            )
        )
    )

    parameters = {
        "document_ids":
            created_document_ids,
    }

    with database_transaction(
        DATABASE_ID,
        ROLE,
    ) as connection:
        connection.execute(
            delete_chunks,
            parameters,
        )

        connection.execute(
            delete_documents,
            parameters,
        )

        remaining_chunks = connection.execute(
            count_chunks,
            parameters,
        ).scalar_one()

        remaining_documents = (
            connection.execute(
                count_documents,
                parameters,
            ).scalar_one()
        )

    return (
        int_value(remaining_documents),
        int_value(remaining_chunks),
    )


def print_results() -> None:
    print()
    print(
        f"{'Test':<42} "
        f"{'Expected':<16} "
        f"{'Actual':<16} "
        "Pass"
    )
    print("-" * 88)

    for item in results:
        print(
            f"{item['name']:<42} "
            f"{str(item['expected']):<16} "
            f"{str(item['actual']):<16} "
            f"{item['pass']}"
        )

        if not item["pass"]:
            print(
                f"  Detail: {item['detail']}"
            )

    passed = sum(
        item["pass"]
        for item in results
    )

    print()
    print(
        f"Result: {passed}/{len(results)} PASS"
    )


def main() -> None:
    settings = get_settings()

    unique = datetime.now().strftime(
        "%Y%m%d%H%M%S%f"
    )

    marker_base = f"LEAFY-{unique}"

    headers = {
        "x-leafy-internal-key":
            settings.leafy_internal_key,
    }

    txt_content = create_txt(
        f"{marker_base}-TXT"
    )

    try:
        with httpx.Client(
            base_url=BASE_URL,
            headers=headers,
            timeout=120,
        ) as client:
            health = client.get("/health")

            add_result(
                "Health",
                200,
                health.status_code,
                health.text,
            )

            with httpx.Client(
                base_url=BASE_URL,
                timeout=30,
            ) as unauthorized_client:
                missing_key = upload_document(
                    unauthorized_client,
                    "missing-key.txt",
                    b"Missing key test",
                    "text/plain",
                )

                add_result(
                    "Missing internal key",
                    401,
                    missing_key.status_code,
                    missing_key.text,
                )

                wrong_key = (
                    unauthorized_client.post(
                        "/documents/ingest",
                        headers={
                            "x-leafy-internal-key":
                                "wrong-key",
                        },
                        data={
                            "role": ROLE,
                            "database_id":
                                DATABASE_ID,
                        },
                        files={
                            "file": (
                                "wrong-key.txt",
                                b"Wrong key test",
                                "text/plain",
                            ),
                        },
                    )
                )

                add_result(
                    "Wrong internal key",
                    401,
                    wrong_key.status_code,
                    wrong_key.text,
                )

            wrong_role = upload_document(
                client,
                "wrong-role.txt",
                b"Wrong role document",
                "text/plain",
                role="user",
            )

            add_result(
                "Wrong role",
                403,
                wrong_role.status_code,
                wrong_role.text,
            )

            unsupported = upload_document(
                client,
                "malware.exe",
                b"MZ fake executable",
                "application/octet-stream",
            )

            add_result(
                "Unsupported extension",
                422,
                unsupported.status_code,
                unsupported.text,
            )

            fake_pdf = upload_document(
                client,
                "fake.pdf",
                b"This is not a real PDF",
                "application/pdf",
            )

            add_result(
                "Invalid PDF signature",
                422,
                fake_pdf.status_code,
                fake_pdf.text,
            )

            empty_file = upload_document(
                client,
                "empty.txt",
                b"",
                "text/plain",
            )

            add_result(
                "Empty document",
                422,
                empty_file.status_code,
                empty_file.text,
            )

            oversized_file = upload_document(
                client,
                "oversized.txt",
                b"A" * (
                    MAX_DOCUMENT_FILE_SIZE + 1
                ),
                "text/plain",
            )

            add_result(
                "Oversized document",
                422,
                oversized_file.status_code,
                oversized_file.text,
            )

            test_valid_upload(
                client=client,
                label="TXT",
                file_name=f"leafy-{unique}.txt",
                file_content=txt_content,
                mime_type="text/plain",
                expected_method="text",
                expected_marker=(
                    f"{marker_base}-TXT"
                ),
            )

            test_valid_upload(
                client=client,
                label="CSV",
                file_name=f"leafy-{unique}.csv",
                file_content=create_csv(
                    f"{marker_base}-CSV"
                ),
                mime_type="text/csv",
                expected_method="csv",
                expected_marker=(
                    f"{marker_base}-CSV"
                ),
            )

            test_valid_upload(
                client=client,
                label="XLSX",
                file_name=f"leafy-{unique}.xlsx",
                file_content=create_xlsx(
                    f"{marker_base}-XLSX"
                ),
                mime_type=(
                    "application/vnd.openxmlformats-"
                    "officedocument.spreadsheetml.sheet"
                ),
                expected_method="xlsx",
                expected_marker=(
                    f"{marker_base}-XLSX"
                ),
            )

            test_valid_upload(
                client=client,
                label="DOCX",
                file_name=f"leafy-{unique}.docx",
                file_content=create_docx(
                    f"{marker_base}-DOCX"
                ),
                mime_type=(
                    "application/vnd.openxmlformats-"
                    "officedocument.wordprocessingml."
                    "document"
                ),
                expected_method="docx",
                expected_marker=(
                    f"{marker_base}-DOCX"
                ),
            )

            test_valid_upload(
                client=client,
                label="PDF text",
                file_name=f"leafy-{unique}.pdf",
                file_content=create_pdf(
                    f"{marker_base}-PDF"
                ),
                mime_type="application/pdf",
                expected_method="pdf_text",
                expected_marker=(
                    f"{marker_base}-PDF"
                ),
            )

            test_valid_upload(
                client=client,
                label="PNG OCR",
                file_name=f"leafy-{unique}.png",
                file_content=create_ocr_image(
                    "LEAFY PNG OCR",
                    "PNG",
                ),
                mime_type="image/png",
                expected_method="image_ocr",
                expected_marker="LEAFY",
            )

            test_valid_upload(
                client=client,
                label="JPG OCR",
                file_name=f"leafy-{unique}.jpg",
                file_content=create_ocr_image(
                    "LEAFY JPG OCR",
                    "JPEG",
                ),
                mime_type="image/jpeg",
                expected_method="image_ocr",
                expected_marker="LEAFY",
            )

            test_valid_upload(
                client=client,
                label="JPEG OCR",
                file_name=f"leafy-{unique}.jpeg",
                file_content=create_ocr_image(
                    "LEAFY JPEG OCR",
                    "JPEG",
                ),
                mime_type="image/jpeg",
                expected_method="image_ocr",
                expected_marker="LEAFY",
            )

            duplicate_response = upload_document(
                client,
                f"duplicate-{unique}.txt",
                txt_content,
                "text/plain",
            )

            add_result(
                "Duplicate content hash",
                409,
                duplicate_response.status_code,
                duplicate_response.text,
            )

            inspect_database()

    finally:
        try:
            (
                remaining_documents,
                remaining_chunks,
            ) = cleanup()

            add_result(
                "Cleanup knowledge_documents",
                0,
                remaining_documents,
            )

            add_result(
                "Cleanup knowledge_chunks",
                0,
                remaining_chunks,
            )

        except Exception as error:
            add_result(
                "Cleanup knowledge_documents",
                0,
                "cleanup_failed",
                repr(error),
            )

            add_result(
                "Cleanup knowledge_chunks",
                0,
                "cleanup_failed",
                repr(error),
            )

        print_results()


if __name__ == "__main__":
    main()